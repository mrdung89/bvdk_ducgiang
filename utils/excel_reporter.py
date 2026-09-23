import os
import openpyxl
from datetime import datetime
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.worksheet.page import PageMargins
from openpyxl.drawing.image import Image as ExcelImage
import subprocess
import platform
from PySide6.QtWidgets import QMessageBox

# Cấu hình đường dẫn lưu trữ mặc định
BASE_REPORT_DIR = os.path.join(os.getcwd(), "BaoCao_Xuat")
LOGO_PATH = os.path.join(os.getcwd(), "assets", "logo.jpg") # Đã trỏ tới folder assets sếp tạo

class ExcelReporter:
    def __init__(self, parent_view=None):
        self.parent_view = parent_view
        
    def _get_safe_name(self, name):
        """Làm sạch chuỗi để làm tên file/thư mục an toàn trên mọi HĐH"""
        if not name: return "Unknown"
        return "".join([c for c in str(name) if c.isalnum() or c in (' ', '_', '-')]).strip()

    def _view_file(self, filepath):
        """Tự động mở file Excel sau khi tạo xong"""
        if not os.path.exists(filepath):
            if self.parent_view:
                QMessageBox.warning(self.parent_view, "Lỗi", "Không tìm thấy file để mở.")
            return

        try:
            if platform.system() == "Windows":
                os.startfile(filepath)
            elif platform.system() == "Darwin":
                subprocess.call(["open", filepath])
            else:
                subprocess.call(["xdg-open", filepath])
        except Exception as e:
            if self.parent_view:
                QMessageBox.warning(self.parent_view, "Lỗi mở file", str(e))

    def _setup_common_styles(self):
        """Khởi tạo các định dạng (font, border) dùng chung"""
        self.font_title = Font(name='Times New Roman', size=16, bold=True)
        self.font_header = Font(name='Times New Roman', size=12, bold=True)
        self.font_normal = Font(name='Times New Roman', size=11)
        self.border_thin = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        self.align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
        self.align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)

    def _draw_common_header(self, ws, report_title, max_col_letter='E'):
        """Vẽ chung Header: Logo bên trái, Tiêu đề bên phải, Tên phiếu ở dưới"""
        # Logo bên trái
        if os.path.exists(LOGO_PATH):
            try:
                img = ExcelImage(LOGO_PATH)
                img.width, img.height = 70, 70
                ws.add_image(img, 'A1')
            except: pass
            
        # Tiêu đề bên phải
        ws.merge_cells(f'B1:{max_col_letter}1'); ws['B1'] = "BỆNH VIỆN ĐA KHOA ĐỨC GIANG"
        ws['B1'].font = self.font_header; ws['B1'].alignment = self.align_center
        ws.merge_cells(f'B2:{max_col_letter}2'); ws['B2'] = "KHOA KIỂM SOÁT NHIỄM KHUẨN"
        ws['B2'].font = self.font_header; ws['B2'].alignment = self.align_center
        
        # Tên Phiếu
        ws.merge_cells(f'A4:{max_col_letter}4'); ws['A4'] = report_title
        ws['A4'].font = self.font_title; ws['A4'].alignment = self.align_center
        ws.row_dimensions[4].height = 25

    # =====================================================================
    # MẪU 1: PHIẾU CẤP PHÁT (Dùng thay thế cho hàm export_issue_csv hiện tại)
    # =====================================================================
    def create_distribution_receipt(self, khoa_nhan, ma_phieu, thoi_gian, data_list, nguoi_lap):
        self._setup_common_styles()
        safe_khoa = self._get_safe_name(khoa_nhan)
        safe_phieu = self._get_safe_name(ma_phieu)
        now = datetime.now()
        folder_path = os.path.join(BASE_REPORT_DIR, "CapPhat", safe_khoa, str(now.year), f"{now.month:02d}", f"{now.day:02d}")
        os.makedirs(folder_path, exist_ok=True)
        filepath = os.path.join(folder_path, f"{safe_phieu}.xlsx")
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = safe_phieu

        ws.page_setup.paperSize = ws.PAPERSIZE_A5
        ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
        ws.sheet_properties.pageSetUpPr.fitToPage = True 
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.2, bottom=0.2)
        
        self._draw_common_header(ws, "PHIẾU GIAO NHẬN / CẤP PHÁT ĐỒ SẠCH", 'E')
        
        ws['A6'] = f"Khoa nhận: {khoa_nhan}"; ws['A6'].font = self.font_header
        ws.merge_cells('D6:E6'); ws['D6'] = f"Mã phiếu: {ma_phieu}"; ws['D6'].alignment = Alignment(horizontal='right')
        ws['A7'] = f"Thời gian: {thoi_gian}"; ws['A7'].font = Font(italic=True)

        headers = ["STT", "Mã Đồ", "Tên Đồ", "Số Lượng", "Ghi Chú"]
        widths = [5, 15, 30, 10, 15]
        start_row = 9
        for col_idx, header in enumerate(headers, 1):
            c = ws.cell(row=start_row, column=col_idx, value=header)
            c.font = self.font_header; c.border = self.border_thin; c.alignment = self.align_center
            c.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = widths[col_idx-1]

        current_row = start_row
        tong_sl = 0
        for i, item in enumerate(data_list, 1):
            current_row += 1
            ws.cell(row=current_row, column=1, value=i).border = self.border_thin
            ws.cell(row=current_row, column=1).alignment = self.align_center
            ws.cell(row=current_row, column=2, value=item.get('ma_do', '')).border = self.border_thin
            ws.cell(row=current_row, column=3, value=item.get('ten_do', '')).border = self.border_thin
            ws.cell(row=current_row, column=3).alignment = self.align_left
            sl = int(item.get('so_luong', 0))
            tong_sl += sl
            c_sl = ws.cell(row=current_row, column=4, value=sl)
            c_sl.border = self.border_thin; c_sl.font = self.font_header; c_sl.alignment = self.align_center
            ws.cell(row=current_row, column=5, value=item.get('ghi_chu', '')).border = self.border_thin

        current_row += 1
        ws.merge_cells(f'A{current_row}:C{current_row}')
        ws[f'A{current_row}'] = "TỔNG CỘNG"
        ws[f'A{current_row}'].font = self.font_header; ws[f'A{current_row}'].alignment = Alignment(horizontal='right')
        ws.cell(row=current_row, column=4, value=tong_sl).font = self.font_header
        
        current_row += 3
        ws.merge_cells(f'A{current_row}:B{current_row}'); ws[f'A{current_row}'] = "NGƯỜI GIAO\n(Ký và ghi rõ họ tên)"
        ws[f'A{current_row}'].font = self.font_header; ws[f'A{current_row}'].alignment = self.align_center
        ws.merge_cells(f'D{current_row}:E{current_row}'); ws[f'D{current_row}'] = "NGƯỜI NHẬN\n(Ký và ghi rõ họ tên)"
        ws[f'D{current_row}'].font = self.font_header; ws[f'D{current_row}'].alignment = self.align_center
        
        current_row += 4
        ws.merge_cells(f'A{current_row}:B{current_row}'); ws[f'A{current_row}'] = nguoi_lap.upper()
        ws[f'A{current_row}'].font = self.font_header; ws[f'A{current_row}'].alignment = self.align_center

        try:
            wb.save(filepath)
            self._view_file(filepath)
            return True
        except Exception as e:
            if self.parent_view: QMessageBox.critical(self.parent_view, "Lỗi Xuất File", str(e))
            return False

    # =====================================================================
    # MẪU 2: NHẬT KÝ HOẠT ĐỘNG
    # =====================================================================
    def create_audit_log_report(self, data_list):
        self._setup_common_styles()
        now = datetime.now()
        folder_path = os.path.join(BASE_REPORT_DIR, "NhatKy", str(now.year), f"{now.month:02d}")
        os.makedirs(folder_path, exist_ok=True)
        filepath = os.path.join(folder_path, f"NhatKy_{now.strftime('%Y%m%d_%H%M%S')}.xlsx")
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "AuditLog"
        
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        
        self._draw_common_header(ws, "NHẬT KÝ HOẠT ĐỘNG HỆ THỐNG", 'E')
        
        headers = ["ID", "Thời Gian", "Người Thực Hiện", "Bảng/Mã", "Nội Dung"]
        widths = [8, 20, 25, 30, 60]
        
        for i, h in enumerate(headers, 1):
            c = ws.cell(row=6, column=i, value=h)
            c.font = self.font_header; c.border = self.border_thin; c.alignment = self.align_center
            c.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = widths[i-1]
            
        for i, item in enumerate(data_list, 1):
            r = 6 + i
            for j, val in enumerate([item['id'], item['thoi_gian'], item['nguoi_thuc_hien'], item['bang_ma'], item['noi_dung']], 1):
                c = ws.cell(row=r, column=j, value=str(val))
                c.border = self.border_thin
                c.alignment = self.align_left if j == 5 else self.align_center

        try:
            wb.save(filepath)
            self._view_file(filepath)
            return True
        except Exception as e:
            if self.parent_view: QMessageBox.critical(self.parent_view, "Lỗi Xuất File", str(e))
            return False

    # =====================================================================
    # MẪU 3: PHIẾU VẬN HÀNH MÁY
    # =====================================================================
    def create_machine_run_report(self, mac_name, cycle_name, start_time, items, nguoi_lap):
        self._setup_common_styles()
        
        # 1. Đường dẫn thư mục: EXCEL VẬN HÀNH / TÊN MÁY / YYYY / MM / DD
        safe_mac = self._get_safe_name(mac_name)
        now = datetime.now()
        folder_path = os.path.join(BASE_REPORT_DIR, "EXCEL VAN HANH", safe_mac, str(now.year), f"{now.month:02d}", f"{now.day:02d}")
        os.makedirs(folder_path, exist_ok=True)
        
        # Tên phiếu: PH_YYMMDDHHMMSS
        ma_phieu = f"PH_{now.strftime('%y%m%d%H%M%S')}"
        filepath = os.path.join(folder_path, f"{ma_phieu}.xlsx")
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = ma_phieu

        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
        ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.2, bottom=0.2)
        
        self._draw_common_header(ws, "PHIẾU VẬN HÀNH MÁY TIỆT KHUẨN", 'E')
        
        ws['A6'] = f"Tên máy: {mac_name}"; ws['A6'].font = self.font_header
        ws.merge_cells('C6:E6'); ws['C6'] = f"Mã mẻ (Phiếu): {ma_phieu}"; ws['C6'].font = self.font_header
        
        ws['A7'] = f"Chu trình: {cycle_name}"; ws['A7'].font = self.font_header
        ws.merge_cells('C7:E7'); ws['C7'] = f"Thời gian bắt đầu: {start_time}"; ws['C7'].font = self.font_header
        
        # Bảng dữ liệu đồ tiệt khuẩn
        headers = ["STT", "Mã đồ / Khoa", "Tên dụng cụ", "Số lượng", "Ghi chú"]
        widths = [5, 20, 35, 10, 20]
        start_row = 9
        for col_idx, header in enumerate(headers, 1):
            c = ws.cell(row=start_row, column=col_idx, value=header)
            c.font = self.font_header; c.border = self.border_thin; c.alignment = self.align_center
            c.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = widths[col_idx-1]

        current_row = start_row
        tong_sl = 0
        for i, item in enumerate(items, 1):
            current_row += 1
            ws.cell(row=current_row, column=1, value=i).border = self.border_thin
            ws.cell(row=current_row, column=1).alignment = self.align_center
            
            ma_khoa = f"{item.get('ma_san_pham', '')}\n({item.get('khoa_giao', '')})"
            c_ma = ws.cell(row=current_row, column=2, value=ma_khoa)
            c_ma.border = self.border_thin; c_ma.alignment = self.align_center
            
            c_ten = ws.cell(row=current_row, column=3, value=item.get('ten_san_pham', ''))
            c_ten.border = self.border_thin; c_ten.alignment = self.align_left
            
            sl = int(item.get('so_luong', 0))
            tong_sl += sl
            c_sl = ws.cell(row=current_row, column=4, value=sl)
            c_sl.border = self.border_thin; c_sl.font = self.font_header; c_sl.alignment = self.align_center
            
            c_ghi = ws.cell(row=current_row, column=5, value="")
            c_ghi.border = self.border_thin

        # Tổng cộng
        current_row += 1
        ws.merge_cells(f'A{current_row}:C{current_row}')
        ws[f'A{current_row}'] = "TỔNG CỘNG"
        ws[f'A{current_row}'].font = self.font_header; ws[f'A{current_row}'].alignment = Alignment(horizontal='right')
        ws.cell(row=current_row, column=4, value=tong_sl).font = self.font_header; ws.cell(row=current_row, column=4).border = self.border_thin
        
        # Form Test kết quả (In sẵn để nv điền tay)
        current_row += 2
        ws.merge_cells(f'A{current_row}:E{current_row}')
        ws[f'A{current_row}'] = "KẾT QUẢ KIỂM TRA (Đánh dấu X hoặc điền thông số):"
        ws[f'A{current_row}'].font = self.font_header
        
        current_row += 1
        ws.merge_cells(f'A{current_row}:E{current_row}')
        ws[f'A{current_row}'] = "[  ] CI Đạt    |    [  ] BI Đạt    |    [  ] Biểu đồ (Nhiệt/Áp suất) Đạt"
        
        current_row += 1
        ws.merge_cells(f'A{current_row}:E{current_row}')
        ws[f'A{current_row}'] = "Kết luận cuối cùng: ........................................................................."
        
        # Chữ ký
        current_row += 3
        ws.merge_cells(f'A{current_row}:B{current_row}'); ws[f'A{current_row}'] = "NGƯỜI VẬN HÀNH\n(Ký và ghi rõ họ tên)"
        ws[f'A{current_row}'].font = self.font_header; ws[f'A{current_row}'].alignment = self.align_center
        
        ws.merge_cells(f'D{current_row}:E{current_row}'); ws[f'D{current_row}'] = "NGƯỜI DỠ HÀNG\n(Ký và ghi rõ họ tên)"
        ws[f'D{current_row}'].font = self.font_header; ws[f'D{current_row}'].alignment = self.align_center
        
        current_row += 4
        ws.merge_cells(f'A{current_row}:B{current_row}'); ws[f'A{current_row}'] = nguoi_lap.upper()
        ws[f'A{current_row}'].font = self.font_header; ws[f'A{current_row}'].alignment = self.align_center

        try:
            wb.save(filepath)
            self._view_file(filepath)
            return True
        except Exception as e:
            if self.parent_view: QMessageBox.critical(self.parent_view, "Lỗi Xuất File", str(e))
            return False
